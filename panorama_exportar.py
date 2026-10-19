#!/usr/bin/env python3
"""Exporta el Panorama de propuestas a archivos para presentar, con los MISMOS datos
que calcula front-panorama.html (consolidado.json + consolidado-unico.json).

Genera en el Escritorio:
  Panorama_propuestas_SOLART.pdf            — 5 láminas como imagen (abre en cualquier lado)
  Panorama_propuestas_SOLART_imagenes.pptx  — las mismas láminas en PowerPoint (imagen por lámina)
  Panorama_propuestas_SOLART.pptx           — versión con gráficas nativas editables (requiere PowerPoint)

Uso:  python3 panorama_exportar.py
"""
import json, math, os, textwrap
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Wedge
from matplotlib.backends.backend_pdf import PdfPages
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION

AQUI = os.path.dirname(os.path.abspath(__file__))
DESK = os.path.expanduser('~/Desktop')
TMP = '/tmp/panorama_qa'; os.makedirs(TMP, exist_ok=True)

# ── Datos (misma lógica que la página) ──────────────────────────────────────
base = {r['id']: r for r in json.load(open(os.path.join(AQUI, 'consolidado.json')))['registros']}
uni = json.load(open(os.path.join(AQUI, 'consolidado-unico.json'))); cl = uni['clientes']
conMat = [c for c in cl if (c.get('nPropuestas') or 0) > 0 or c.get('excel')]
soloCarp = [c for c in cl if c.get('carpeta') and c not in conMat]
nProp = sum(c.get('nPropuestas') or 0 for c in conMat)
est = {'abierto': 0, 'sinCRM': 0, 'ganado': 0, 'perdido': 0}; dist = {0: 0, 1: 0, 2: 0, 3: 0, 4: 0}
usd = kw = 0; kwN = 0
for c in conMat:
    r = base.get(c.get('idCRM')); est[r['estado'] if r else 'sinCRM'] += 1
    dist[min(c.get('nPropuestas') or 0, 4)] += 1
    props = sorted([p for p in (c.get('propuestas') or []) if p.get('precioUSD')], key=lambda p: str(p.get('fecha') or ''), reverse=True)
    porSitio = {}
    for p in props:
        k = str(p.get('carpeta') or ''); porSitio.setdefault(k.split(' / ')[0].strip() if ' / ' in k else '', p)
    usd += sum(p['precioUSD'] for p in porSitio.values())
    kws = [p['kwp'] for p in porSitio.values() if p.get('kwp')]
    if kws: kw += sum(kws); kwN += 1
total = len(conMat); vivos = est['abierto'] + est['sinCRM']; rep = dist[2] + dist[3] + dist[4]
gen = str(uni.get('generado', ''))[:10]
fmt = lambda n: f'{round(n):,}'

# ── Paleta ──────────────────────────────────────────────────────────────────
NAVY, CARD, LINE = '#0B1B2B', '#112A44', '#1E3A5F'
GOLD, SKY, GREEN, GRAY, WHITE, MUTED, SOFT, INK = '#F5C042', '#7DD3FC', '#4ADE80', '#64748B', '#FFFFFF', '#94A3B8', '#CBD5E1', '#1A1A2E'
W, H = 13.333, 7.5
plt.rcParams['font.family'] = ['Helvetica Neue', 'Helvetica', 'Arial', 'DejaVu Sans']

def lienzo():
    fig = plt.figure(figsize=(W, H), dpi=160); fig.patch.set_facecolor(NAVY)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, W); ax.set_ylim(0, H); ax.axis('off'); return fig, ax
def card(ax, x, y, w, h, fc=CARD, ec=LINE):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0,rounding_size=0.22', fc=fc, ec=ec, lw=1.2))
def T(ax, x, y, s, size=16, color=WHITE, bold=False, ha='left', va='top', italic=False):
    ax.text(x, y, s, fontsize=size, color=color, fontweight='bold' if bold else 'normal', ha=ha, va=va,
            fontstyle='italic' if italic else 'normal', linespacing=1.25)
def titulo(ax, s): T(ax, 0.6, H - 0.5, s, 30, WHITE, True)

segs = [('En proceso', 'con negociación abierta en el CRM', est['abierto'], SKY),
        ('Sin seguimiento', 'tienen propuesta, pero no están en el CRM', est['sinCRM'], GOLD),
        ('Ganados', 'proyectos cerrados', est['ganado'], GREEN),
        ('Perdidos', 'cerrados sin venta', est['perdido'], GRAY)]
cajas = [(fmt(total + len(soloCarp)), 'clientes a los que les\nhemos propuesto', f'{total} con propuesta leída\n{len(soloCarp)} con carpeta sin descargar', WHITE),
         (fmt(nProp), 'propuestas elaboradas', f'{nProp/total:.1f} por cliente en promedio', SKY),
         (f'${usd/1e6:.1f} M', 'USD en propuestas\nvigentes', 'última propuesta de cada sitio,\nsin duplicar versiones', GOLD),
         (fmt(kw), 'kW propuestos', f'en los {kwN} clientes\ncon tamaño legible', GREEN)]
pasos = [f'Retomar los {est["sinCRM"]} clientes con propuesta que no están en el CRM: darlos de alta y asignar vendedor.',
         f'Revisar los {est["abierto"]} en proceso: cuáles tienen propuesta vigente y cuál es su siguiente paso.',
         f'Descargar las {len(soloCarp)} carpetas pendientes de OneDrive para que entren al panorama.']
mensaje1 = f'{vivos} clientes con propuesta siguen vivos: {est["abierto"]} en negociación\ny {est["sinCRM"]} a los que hay que retomar.'
mensaje2 = f'Solo hemos ganado {est["ganado"]} de {total}. El potencial está sobre la mesa.'
mensaje3 = f'Si cerráramos 1 de cada 10 vivos, serían {round(vivos/10)} proyectos más.'
nota = ('Cómo se calcula: el universo son los clientes con al menos una propuesta leída en OneDrive. El estado sale del CRM '
        '(Bitrix o Pipedrive); "sin seguimiento" son propuestas sin registro ahí. El valor suma la propuesta más reciente de cada '
        'sitio, sin contar versiones anteriores. Las carpetas sin descargar no entran al pastel.')

# ── Láminas como imagen ─────────────────────────────────────────────────────
figs = []
fig, ax = lienzo()
T(ax, 0.9, 5.1, 'Panorama de propuestas', 52, WHITE, True)
T(ax, 0.9, 4.15, 'Todo lo que SOLART ha propuesto, en una sola gráfica', 24, GOLD)
T(ax, 0.9, 3.4, f'{total} clientes con propuesta · {fmt(nProp)} propuestas · datos al {gen}', 16, MUTED)
T(ax, 0.9, 0.7, 'Fuente: carpetas de propuestas en OneDrive, cruzadas con el CRM (Bitrix y Pipedrive)', 11, MUTED, italic=True)
figs.append(fig)

fig, ax = lienzo(); titulo(ax, 'El tamaño de lo que ya hicimos')
for i, (v, l, sub, col) in enumerate(cajas):
    x = 0.6 + i * 3.1; card(ax, x, 2.1, 2.9, 3.7)
    T(ax, x + 0.3, 5.2, v, 40, col, True); T(ax, x + 0.3, 4.05, l, 15, SOFT, True); T(ax, x + 0.3, 3.1, sub, 11, MUTED)
T(ax, 0.6, 1.5, 'Cada propuesta es tiempo de ingeniería y de ventas ya invertido.\nLa pregunta de esta reunión es cuánto de esto podemos convertir.', 16, SOFT, italic=True)
figs.append(fig)

fig, ax = lienzo(); titulo(ax, 'Nuestro mercado: ¿en qué está cada cliente con propuesta?')
cx, cy, R, r = 3.6, 3.55, 2.35, 1.45; ang = 90
for t, d, n, col in segs:
    sw = n / total * 360
    ax.add_patch(Wedge((cx, cy), R, ang - sw, ang, width=R - r, fc=col, ec=NAVY, lw=3))
    mid = math.radians(ang - sw / 2)
    if n / total > 0.05:
        ax.text(cx + (R + r) / 2 * math.cos(mid), cy + (R + r) / 2 * math.sin(mid), str(n), ha='center', va='center', fontsize=18, fontweight='bold', color=INK if col != GRAY else WHITE)
    ang -= sw
T(ax, cx, cy + 0.15, str(total), 46, WHITE, True, 'center', 'center'); T(ax, cx, cy - 0.55, 'clientes con propuesta', 12, SOFT, ha='center', va='center')
card(ax, 7.3, 1.0, 5.4, 5.35)
for i, (t, d, n, col) in enumerate(segs):
    y = 6.0 - i * 1.3
    ax.add_patch(FancyBboxPatch((7.65, y - 0.42), 0.3, 0.3, boxstyle='round,pad=0,rounding_size=0.06', fc=col, ec='none'))
    T(ax, 8.15, y, t, 18, WHITE, True); T(ax, 8.15, y - 0.5, d, 11, MUTED)
    T(ax, 12.4, y, str(n), 24, col, True, ha='right'); T(ax, 12.4, y - 0.52, f'{round(n/total*100)}%', 12, MUTED, ha='right')
figs.append(fig)

fig, ax = lienzo(); titulo(ax, 'Insistimos: clientes por número de propuestas')
cats = [('1 propuesta', dist[1]), ('2 propuestas', dist[2]), ('3 propuestas', dist[3]), ('4 o más', dist[4]), ('Solo hoja de cálculo', dist[0])]
mx = max(v for _, v in cats); x0, bw = 3.0, 4.9
for i, (l, v) in enumerate(cats):
    y = 5.8 - i * 1.0
    T(ax, 2.8, y + 0.3, l, 15, SOFT, ha='right', va='center')
    ax.add_patch(FancyBboxPatch((x0, y), bw, 0.6, boxstyle='round,pad=0,rounding_size=0.12', fc='#0F2238', ec='none'))
    ax.add_patch(FancyBboxPatch((x0, y), max(bw * v / mx, 0.5), 0.6, boxstyle='round,pad=0,rounding_size=0.12', fc=GOLD if l == '4 o más' else SKY, ec='none'))
    T(ax, x0 + max(bw * v / mx, 0.5) - 0.15, y + 0.3, str(v), 15, INK, True, ha='right', va='center')
card(ax, 8.6, 1.0, 4.1, 5.35)
T(ax, 8.95, 6.1, str(rep), 54, GOLD, True); T(ax, 8.95, 4.8, f'clientes ({round(rep/total*100)}%) ya\nrecibieron dos o más\npropuestas', 16, WHITE, True)
T(ax, 8.95, 3.3, 'Nadie pide una tercera propuesta\nsi no le interesa. Ahí hay interés\nreal; lo que falta es cerrar.', 13, SOFT, italic=True)
T(ax, 8.95, 1.8, f'{dist[4]} clientes van por la cuarta\npropuesta o más.', 12, MUTED)
figs.append(fig)

fig, ax = lienzo()
card(ax, 0.9, 2.6, 11.5, 3.6, GOLD, GOLD)
T(ax, 1.4, 5.8, mensaje1, 27, INK, True); T(ax, 1.4, 4.0, mensaje2, 20, INK); T(ax, 1.4, 3.3, mensaje3, 15, INK, italic=True)
T(ax, 0.9, 2.1, 'Siguientes pasos', 16, GOLD, True)
for i, p in enumerate(pasos): T(ax, 1.1, 1.6 - i * 0.42, f'{i+1}.  {p}', 14, SOFT)
figs.append(fig)

pngs = []
for i, f in enumerate(figs, 1):
    p = f'{TMP}/slide-{i}.png'; f.savefig(p, dpi=160, facecolor=NAVY); pngs.append(p)
with PdfPages(f'{DESK}/Panorama_propuestas_SOLART.pdf') as pdf:
    for f in figs: pdf.savefig(f, facecolor=NAVY)

prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)
for p in pngs:
    s = prs.slides.add_slide(prs.slide_layouts[6]); s.shapes.add_picture(p, 0, 0, prs.slide_width, prs.slide_height)
    s.notes_slide.notes_text_frame.text = nota
prs.save(f'{DESK}/Panorama_propuestas_SOLART_imagenes.pptx')

# ── Versión con gráficas nativas (editable en PowerPoint) ───────────────────
def rgb(h): return RGBColor.from_string(h.lstrip('#'))
prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H); BL = prs.slide_layouts[6]
def fondo(s): f = s.background.fill; f.solid(); f.fore_color.rgb = rgb(NAVY)
def texto(s, x, y, w, h, t, size=16, bold=False, color=WHITE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align; r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = rgb(color); r.font.name = 'Calibri'; r.font.italic = italic
def tarjeta(s, x, y, w, h, fill=CARD):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); sh.adjustments[0] = 0.08
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill); sh.line.color.rgb = rgb(LINE); sh.line.width = Pt(0.75); sh.shadow.inherit = False
def tit(s, t): texto(s, 0.6, 0.45, 12.1, 0.8, t, 34, True)

s = prs.slides.add_slide(BL); fondo(s)
texto(s, 0.9, 2.2, 11.5, 1.3, 'Panorama de propuestas', 54, True); texto(s, 0.9, 3.45, 11.5, 0.7, 'Todo lo que SOLART ha propuesto, en una sola gráfica', 24, color=GOLD)
texto(s, 0.9, 4.3, 11.5, 0.5, f'{total} clientes con propuesta · {fmt(nProp)} propuestas · datos al {gen}', 16, color=MUTED)
s = prs.slides.add_slide(BL); fondo(s); tit(s, 'El tamaño de lo que ya hicimos')
for i, (v, l, sub, col) in enumerate(cajas):
    x = 0.6 + i * 3.1; tarjeta(s, x, 1.7, 2.9, 3.6)
    texto(s, x + 0.3, 2.0, 2.3, 1.1, v, 44, True, col, anchor=MSO_ANCHOR.BOTTOM); texto(s, x + 0.3, 3.2, 2.3, 0.8, l.replace('\n', ' '), 15, True, SOFT); texto(s, x + 0.3, 4.1, 2.3, 1.0, sub.replace('\n', ' '), 11, color=MUTED)
s = prs.slides.add_slide(BL); fondo(s); tit(s, 'Nuestro mercado: ¿en qué está cada cliente con propuesta?')
cd = CategoryChartData(); cd.categories = [t for t, _, _, _ in segs]; cd.add_series('Clientes', tuple(n for _, _, n, _ in segs))
ch = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(0.6), Inches(1.4), Inches(6.2), Inches(5.6), cd).chart
ch.has_legend = False; pl = ch.plots[0]; pl.has_data_labels = True; pl.data_labels.font.size = Pt(16); pl.data_labels.font.bold = True; pl.data_labels.font.color.rgb = rgb(INK)
for idx, (_, _, _, col) in enumerate(segs):
    pt = pl.series[0].points[idx]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(col)
tarjeta(s, 7.3, 1.5, 5.4, 5.4)
for i, (t, d, n, col) in enumerate(segs):
    y = 1.85 + i * 1.3; texto(s, 8.15, y, 3.0, 0.45, t, 18, True); texto(s, 8.15, y + 0.48, 3.1, 0.6, d, 11, color=MUTED)
    texto(s, 11.1, y, 1.3, 0.45, str(n), 24, True, col, PP_ALIGN.RIGHT); texto(s, 11.1, y + 0.5, 1.3, 0.4, f'{round(n/total*100)}%', 12, color=MUTED, align=PP_ALIGN.RIGHT)
s = prs.slides.add_slide(BL); fondo(s); tit(s, 'Insistimos: clientes por número de propuestas')
cd = CategoryChartData(); cd.categories = [l for l, _ in reversed(cats)]; cd.add_series('Clientes', tuple(v for _, v in reversed(cats)))
ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.4), Inches(7.6), Inches(5.4), cd).chart
ch.has_legend = False; pl = ch.plots[0]; pl.gap_width = 60; pl.has_data_labels = True; pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
pl.data_labels.font.size = Pt(14); pl.data_labels.font.bold = True; pl.data_labels.font.color.rgb = rgb(WHITE)
ser = pl.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(SKY)
ch.category_axis.tick_labels.font.size = Pt(13); ch.category_axis.tick_labels.font.color.rgb = rgb(SOFT); ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False
tarjeta(s, 8.6, 1.5, 4.1, 5.3); texto(s, 8.95, 1.9, 3.4, 1.1, str(rep), 54, True, GOLD)
texto(s, 8.95, 3.0, 3.4, 0.9, f'clientes ({round(rep/total*100)}%) ya recibieron dos o más propuestas', 16, True)
texto(s, 8.95, 4.1, 3.4, 1.6, 'Nadie pide una tercera propuesta si no le interesa. Ahí hay interés real; lo que falta es cerrar.', 14, color=SOFT, italic=True)
s = prs.slides.add_slide(BL); fondo(s); tarjeta(s, 0.9, 1.4, 11.5, 3.6, GOLD)
texto(s, 1.4, 1.8, 10.5, 2.0, mensaje1.replace('\n', ' '), 28, True, INK); texto(s, 1.4, 3.65, 10.5, 0.6, mensaje2, 20, color=INK); texto(s, 1.4, 4.3, 10.5, 0.5, mensaje3, 15, color=INK, italic=True)
texto(s, 0.9, 5.3, 11.5, 0.5, 'Siguientes pasos', 16, True, GOLD)
for i, p in enumerate(pasos): texto(s, 1.1, 5.8 + i * 0.42, 11.2, 0.4, f'{i+1}.  {p}', 14, color=SOFT)
for sl in prs.slides: sl.notes_slide.notes_text_frame.text = nota
prs.save(f'{DESK}/Panorama_propuestas_SOLART.pptx')

print(f'Listo · {total} clientes · {fmt(nProp)} propuestas · ${usd/1e6:.1f} M USD · {fmt(kw)} kW · estados {est}')
print('Imágenes de revisión en', TMP)
